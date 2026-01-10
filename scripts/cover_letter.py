# scripts/cover_letter.py
"""Génération de lettre de motivation depuis content.json."""

import json
from pathlib import Path
from datetime import datetime
from docx import Document

from utils.config import TEMPLATES, OUTPUT_NAMES, COVER_LETTER_PLACEHOLDERS
from utils.docx import replace_all_placeholders, remove_empty_paragraphs
from utils.pdf import convert_and_trim


def generate_cover_letter(content: dict, output_folder: str | Path) -> tuple[Path, Path | None]:
    """
    Génère une lettre de motivation depuis le content.json.

    Args:
        content: Dict contenant les données (depuis content.json)
        output_folder: Dossier de destination

    Returns:
        Tuple (docx_path, pdf_path)
    """
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    # Charger le template
    doc = Document(TEMPLATES["cover_letter"])

    # Extraire les données
    metadata = content.get("metadata", {})
    cl_data = content.get("cover_letter", {})

    # Date et salutation
    current_date = datetime.now().strftime("%B %d, %Y")
    recipient = cl_data.get("recipient", "Hiring Manager")
    salutation = f"Dear {recipient},"

    # Construire les remplacements
    replacements = {
        COVER_LETTER_PLACEHOLDERS["date"]: current_date,
        COVER_LETTER_PLACEHOLDERS["recipient"]: recipient,
        COVER_LETTER_PLACEHOLDERS["street"]: cl_data.get("street", ""),
        COVER_LETTER_PLACEHOLDERS["postal"]: cl_data.get("postal", ""),
        COVER_LETTER_PLACEHOLDERS["company"]: metadata.get("company", ""),
        COVER_LETTER_PLACEHOLDERS["salutation"]: salutation,
        COVER_LETTER_PLACEHOLDERS["intro"]: cl_data.get("intro", ""),
        COVER_LETTER_PLACEHOLDERS["body_1"]: cl_data.get("body_1", ""),
        COVER_LETTER_PLACEHOLDERS["body_2"]: cl_data.get("body_2", ""),
        COVER_LETTER_PLACEHOLDERS["body_3"]: cl_data.get("body_3", ""),
        COVER_LETTER_PLACEHOLDERS["additional"]: cl_data.get("additional", ""),
        COVER_LETTER_PLACEHOLDERS["attraction"]: cl_data.get("attraction", ""),
        COVER_LETTER_PLACEHOLDERS["closing"]: cl_data.get("closing", "")
    }

    # Identifier les placeholders optionnels vides à supprimer
    placeholders_to_remove = []
    if not cl_data.get("additional"):
        placeholders_to_remove.append(COVER_LETTER_PLACEHOLDERS["additional"])
    if not cl_data.get("attraction"):
        placeholders_to_remove.append(COVER_LETTER_PLACEHOLDERS["attraction"])

    # Supprimer les paragraphes optionnels vides
    remove_empty_paragraphs(doc, placeholders_to_remove)

    # Remplacer les placeholders (avec support markdown pour **bold**)
    replace_all_placeholders(doc, replacements, support_markdown=True)

    # Supprimer les paragraphes vides trailing
    remove_empty_paragraphs(doc)

    # Sauvegarder le DOCX
    docx_path = output_folder / OUTPUT_NAMES["cover_letter_docx"]
    doc.save(docx_path)
    print(f"[DOCX] {docx_path}")

    # Générer le PDF
    pdf_folder = output_folder / "PDF"
    pdf_path = convert_and_trim(docx_path, pdf_folder)
    if pdf_path:
        print(f"[PDF] {pdf_path} (1 page)")

    return docx_path, pdf_path


def generate_cover_letter_from_file(content_path: str | Path, output_folder: str | Path = None) -> tuple[Path, Path | None]:
    """
    Génère une lettre depuis un fichier content.json.

    Args:
        content_path: Chemin vers content.json
        output_folder: Dossier de destination (défaut: même dossier que content.json)

    Returns:
        Tuple (docx_path, pdf_path)
    """
    content_path = Path(content_path)

    with open(content_path, 'r', encoding='utf-8') as f:
        content = json.load(f)

    if output_folder is None:
        output_folder = content_path.parent

    return generate_cover_letter(content, output_folder)
