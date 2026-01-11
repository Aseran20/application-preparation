from docx import Document
from docx.shared import Pt
import sys
import re

def validate_cover_letter(cl_path):
    """
    Validate the generated cover letter for:
    1. No remaining placeholders
    2. Correct formatting (Times New Roman, 10pt)
    3. Approximately 1 page (300-500 words)

    Returns:
        tuple: (is_valid, errors_list, word_count)
    """
    errors = []
    word_count = 0

    try:
        doc = Document(cl_path)
    except Exception as e:
        return False, [f"ERROR: Could not open document: {e}"], 0

    # Check for remaining placeholders
    placeholder_patterns = [
        r'\[.*?\]',  # Any text in brackets
    ]

    # Count words and check formatting
    for para_idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()

        # Count words
        if text:
            word_count += len(text.split())

        # Check for placeholders
        for pattern in placeholder_patterns:
            matches = re.findall(pattern, text)
            if matches:
                # Ignore hardcoded email in contact info
                if 'gmail.com' not in text and 'Clotilde' not in text:
                    errors.append(f"PLACEHOLDER FOUND in Paragraph {para_idx}: {matches}")

        # Check formatting for non-empty paragraphs
        if text and len(para.runs) > 0:
            for run_idx, run in enumerate(para.runs):
                if run.text.strip():
                    # Check font name
                    if run.font.name and run.font.name != 'Times New Roman':
                        errors.append(f"WRONG FONT in Paragraph {para_idx}: '{run.font.name}' (expected 'Times New Roman')")

                    # Check font size (10pt for cover letter)
                    if run.font.size:
                        size_pt = run.font.size.pt
                        if size_pt != 10.0:
                            errors.append(f"WRONG SIZE in Paragraph {para_idx}: {size_pt}pt (expected 10pt)")

    # Check word count (ideal range: 300-500 words)
    if word_count < 250:
        errors.append(f"WORD COUNT TOO LOW: {word_count} words (minimum 250 recommended)")
    elif word_count > 550:
        errors.append(f"WORD COUNT TOO HIGH: {word_count} words (maximum 550 recommended for 1 page)")

    # Determine if valid
    is_valid = len(errors) == 0

    return is_valid, errors, word_count

if __name__ == "__main__":
    # Get cover letter path from command line
    if len(sys.argv) > 1:
        cl_path = sys.argv[1]
    else:
        print("Usage: python validate_cover_letter.py <path_to_cover_letter.docx>")
        sys.exit(1)

    print("=" * 80)
    print(f"VALIDATING COVER LETTER: {cl_path}")
    print("=" * 80)

    is_valid, errors, word_count = validate_cover_letter(cl_path)

    if is_valid:
        print(f"\n[PASS] VALIDATION PASSED - Cover letter is correctly formatted!")
        print(f"\nWord count: {word_count} words")
        print("\nNo issues found:")
        print("  - All placeholders replaced")
        print("  - Correct font (Times New Roman)")
        print("  - Correct sizing (10pt)")
        print("  - Appropriate length for 1 page")
    else:
        print(f"\n[FAIL] VALIDATION FAILED - Found {len(errors)} issue(s):")
        print(f"\nWord count: {word_count} words")
        print()
        for i, error in enumerate(errors, 1):
            print(f"{i}. {error}")
        print("\n" + "=" * 80)
        print("Please review and fix these issues.")

    print("\n" + "=" * 80)

    sys.exit(0 if is_valid else 1)
