# generator/scripts/resume.py
"""Génération de CV depuis content.json."""

import json
import sys
from pathlib import Path
from docx import Document

# Add project root to path for imports
sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from generator.utils.config import TEMPLATES, get_output_names, RESUME_PLACEHOLDERS
from generator.utils.docx import replace_all_placeholders, remove_empty_paragraphs
from generator.utils.pdf import convert_and_trim


def generate_resume(content: dict, output_folder: str | Path) -> tuple[Path, Path | None]:
    """
    Génère un CV depuis le content.json.

    Args:
        content: Dict contenant les données (depuis content.json)
        output_folder: Dossier de destination

    Returns:
        Tuple (docx_path, pdf_path)
    """
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    # Récupérer le nom de l'entreprise pour les noms de fichiers
    company = content.get("metadata", {}).get("company", "Company")
    output_names = get_output_names(company)

    # Charger le template
    doc = Document(TEMPLATES["resume"])

    # Construire les remplacements depuis content.json
    resume_data = content.get("resume", {})

    replacements = {
        RESUME_PLACEHOLDERS["introduction"]: resume_data.get("professional_summary", ""),
        RESUME_PLACEHOLDERS["auraia_1"]: resume_data.get("auraia_bullets", ["", "", ""])[0],
        RESUME_PLACEHOLDERS["auraia_2"]: resume_data.get("auraia_bullets", ["", "", ""])[1],
        RESUME_PLACEHOLDERS["auraia_3"]: resume_data.get("auraia_bullets", ["", "", ""])[2],
        RESUME_PLACEHOLDERS["rc"]: resume_data.get("rc_bullet", ""),
        RESUME_PLACEHOLDERS["europ"]: resume_data.get("europ_bullet", ""),
        RESUME_PLACEHOLDERS["leadership_1"]: resume_data.get("leadership_bullets", ["", "", ""])[0],
        RESUME_PLACEHOLDERS["leadership_2"]: resume_data.get("leadership_bullets", ["", "", ""])[1],
        RESUME_PLACEHOLDERS["leadership_3"]: resume_data.get("leadership_bullets", ["", "", ""])[2],
        RESUME_PLACEHOLDERS["courses"]: resume_data.get("courses", ""),
        RESUME_PLACEHOLDERS["skills"]: resume_data.get("skills", "")
    }

    # Remplacer les placeholders
    replace_all_placeholders(doc, replacements, support_markdown=False)

    # Supprimer les paragraphes vides trailing
    remove_empty_paragraphs(doc)

    # Sauvegarder le DOCX
    docx_path = output_folder / output_names["resume_docx"]
    doc.save(docx_path)
    print(f"[DOCX] {docx_path}")

    # Générer le PDF
    pdf_folder = output_folder / "PDF"
    pdf_path, original_pages = convert_and_trim(docx_path, pdf_folder)
    if pdf_path:
        if original_pages > 1:
            print(f"[WARNING] Resume exceeded 1 page ({original_pages} pages) - content was trimmed!")
            print(f"[WARNING] Please shorten the content in content.json and regenerate.")
        print(f"[PDF] {pdf_path} (trimmed to 1 page)")

    return docx_path, pdf_path


def generate_resume_from_file(content_path: str | Path, output_folder: str | Path = None) -> tuple[Path, Path | None]:
    """
    Génère un CV depuis un fichier content.json.

    Args:
        content_path: Chemin vers content.json
        output_folder: Dossier de destination (défaut: même dossier que content.json)

    Returns:
        Tuple (docx_path, pdf_path)
    """
    content_path = Path(content_path)

    with open(content_path, 'r', encoding='utf-8') as f:
        content = json.load(f)

    if output_folder is None:
        output_folder = content_path.parent

    return generate_resume(content, output_folder)
