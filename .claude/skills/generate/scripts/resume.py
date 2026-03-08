# generator/scripts/resume.py
"""Resume generation from content.json."""

import json
import sys
from pathlib import Path
from docx import Document

sys.path.insert(0, str(Path(__file__).parent))

from utils.config import TEMPLATES, get_output_names, RESUME_PLACEHOLDERS
from utils.docx import replace_all_placeholders, remove_empty_paragraphs
from utils.pdf import convert_and_trim


def generate_resume(content: dict, output_folder: str | Path) -> tuple[Path, Path | None]:
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    company = content.get("metadata", {}).get("company", "Company")
    output_names = get_output_names(company)

    doc = Document(TEMPLATES["resume"])
    resume_data = content.get("resume", {})

    auraia   = resume_data.get("auraia_bullets",   [""] * 4)
    rc       = resume_data.get("rc_bullets",       [""] * 3)
    generali = resume_data.get("generali_bullets", [""] * 3)
    extra    = resume_data.get("extra_entries",    [{"title": "", "sentence": ""}] * 3)

    replacements = {
        RESUME_PLACEHOLDERS["introduction"]:         resume_data.get("introduction", ""),
        RESUME_PLACEHOLDERS["auraia_1"]:             auraia[0]   if len(auraia)   > 0 else "",
        RESUME_PLACEHOLDERS["auraia_2"]:             auraia[1]   if len(auraia)   > 1 else "",
        RESUME_PLACEHOLDERS["auraia_3"]:             auraia[2]   if len(auraia)   > 2 else "",
        RESUME_PLACEHOLDERS["auraia_4"]:             auraia[3]   if len(auraia)   > 3 else "",
        RESUME_PLACEHOLDERS["rc_1"]:                 rc[0]       if len(rc)       > 0 else "",
        RESUME_PLACEHOLDERS["rc_2"]:                 rc[1]       if len(rc)       > 1 else "",
        RESUME_PLACEHOLDERS["rc_3"]:                 rc[2]       if len(rc)       > 2 else "",
        RESUME_PLACEHOLDERS["generali_1"]:           generali[0] if len(generali) > 0 else "",
        RESUME_PLACEHOLDERS["generali_2"]:           generali[1] if len(generali) > 1 else "",
        RESUME_PLACEHOLDERS["generali_3"]:           generali[2] if len(generali) > 2 else "",
        RESUME_PLACEHOLDERS["extra_category_title"]: resume_data.get("extra_category_title", "ENTREPRENEURIAL & INDEPENDENT PROJECTS"),
        RESUME_PLACEHOLDERS["extra_bp1_title"]:      extra[0].get("title", "")    if len(extra) > 0 else "",
        RESUME_PLACEHOLDERS["extra_bp1_content"]:    extra[0].get("sentence", "") if len(extra) > 0 else "",
        RESUME_PLACEHOLDERS["extra_bp2_title"]:      extra[1].get("title", "")    if len(extra) > 1 else "",
        RESUME_PLACEHOLDERS["extra_bp2_content"]:    extra[1].get("sentence", "") if len(extra) > 1 else "",
        RESUME_PLACEHOLDERS["extra_bp3_title"]:      extra[2].get("title", "")    if len(extra) > 2 else "",
        RESUME_PLACEHOLDERS["extra_bp3_content"]:    extra[2].get("sentence", "") if len(extra) > 2 else "",
        RESUME_PLACEHOLDERS["extra_bp4_title"]:      extra[3].get("title", "")    if len(extra) > 3 else "",
        RESUME_PLACEHOLDERS["extra_bp4_content"]:    extra[3].get("sentence", "") if len(extra) > 3 else "",
        RESUME_PLACEHOLDERS["extra_bp5_title"]:      extra[4].get("title", "")    if len(extra) > 4 else "",
        RESUME_PLACEHOLDERS["extra_bp5_content"]:    extra[4].get("sentence", "") if len(extra) > 4 else "",
        RESUME_PLACEHOLDERS["coursework"]:           resume_data.get("coursework", ""),
        RESUME_PLACEHOLDERS["tools"]:                resume_data.get("tools", ""),
        RESUME_PLACEHOLDERS["skills"]:               resume_data.get("skills", ""),
        RESUME_PLACEHOLDERS["interests"]:            resume_data.get("interests", ""),
    }

    # Collect unused extra entry placeholders to remove their paragraphs
    unused_placeholders = []
    for i in range(len(extra), 5):
        unused_placeholders.append(RESUME_PLACEHOLDERS[f"extra_bp{i+1}_title"])

    # Title bold formatting is native in the Word template, no markdown parsing needed
    replace_all_placeholders(doc, replacements, support_markdown=False)
    remove_empty_paragraphs(doc, unused_placeholders)

    docx_path = output_folder / output_names["resume_docx"]
    doc.save(docx_path)
    print(f"[DOCX] {docx_path}")

    pdf_folder = output_folder / "PDF"
    pdf_path, original_pages = convert_and_trim(docx_path, pdf_folder)
    if pdf_path:
        if original_pages > 1:
            print(f"[WARNING] Resume exceeded 1 page ({original_pages} pages) - trimmed!")
        print(f"[PDF] {pdf_path}")

    return docx_path, pdf_path


def generate_resume_from_file(content_path: str | Path, output_folder: str | Path = None) -> tuple[Path, Path | None]:
    content_path = Path(content_path)
    with open(content_path, 'r', encoding='utf-8') as f:
        content = json.load(f)
    if output_folder is None:
        output_folder = content_path.parent
    return generate_resume(content, output_folder)
