# generator/scripts/cover_letter.py
"""Cover letter generation from content.json."""

import json
import sys
from pathlib import Path
from datetime import datetime
from docx import Document

sys.path.insert(0, str(Path(__file__).parent))

from utils.config import TEMPLATES, get_output_names, COVER_LETTER_PLACEHOLDERS
from utils.docx import replace_all_placeholders, remove_empty_paragraphs
from utils.pdf import convert_and_trim


def generate_cover_letter(content: dict, output_folder: str | Path) -> tuple[Path, Path | None]:
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    metadata = content.get("metadata", {})
    cl_data  = content.get("cover_letter", {})

    company      = metadata.get("company", "Company")
    output_names = get_output_names(company)

    doc = Document(TEMPLATES["cover_letter"])

    current_date = datetime.now().strftime("%B %d, %Y")
    recipient    = cl_data.get("recipient_name", "Hiring Manager")
    salutation   = cl_data.get("salutation", f"Dear {recipient},")

    replacements = {
        COVER_LETTER_PLACEHOLDERS["date"]:                current_date,
        COVER_LETTER_PLACEHOLDERS["recipient_name"]:      recipient,
        COVER_LETTER_PLACEHOLDERS["street_number"]:       cl_data.get("street_number", ""),
        COVER_LETTER_PLACEHOLDERS["postal_city_country"]: cl_data.get("postal_city_country", ""),
        COVER_LETTER_PLACEHOLDERS["company_name"]:        cl_data.get("company_name", metadata.get("company", "")),
        COVER_LETTER_PLACEHOLDERS["subject_line"]:        cl_data.get("subject_line", ""),
        COVER_LETTER_PLACEHOLDERS["salutation"]:          salutation,
        COVER_LETTER_PLACEHOLDERS["paragraph_1"]:         cl_data.get("paragraph_1", ""),
        COVER_LETTER_PLACEHOLDERS["paragraph_2"]:         cl_data.get("paragraph_2", ""),
        COVER_LETTER_PLACEHOLDERS["paragraph_3"]:         cl_data.get("paragraph_3", ""),
        COVER_LETTER_PLACEHOLDERS["closing"]:             cl_data.get("closing", ""),
    }

    # Remove optional address lines if empty
    placeholders_to_remove = []
    if not cl_data.get("street_number"):
        placeholders_to_remove.append(COVER_LETTER_PLACEHOLDERS["street_number"])
    if not cl_data.get("postal_city_country"):
        placeholders_to_remove.append(COVER_LETTER_PLACEHOLDERS["postal_city_country"])
    remove_empty_paragraphs(doc, placeholders_to_remove)

    replace_all_placeholders(doc, replacements, support_markdown=False)
    remove_empty_paragraphs(doc)

    docx_path = output_folder / output_names["cover_letter_docx"]
    doc.save(docx_path)
    print(f"[DOCX] {docx_path}")

    pdf_folder = output_folder / "PDF"
    pdf_path, original_pages = convert_and_trim(docx_path, pdf_folder)
    if pdf_path:
        if original_pages > 1:
            print(f"[WARNING] Cover letter exceeded 1 page ({original_pages} pages) - trimmed!")
        print(f"[PDF] {pdf_path}")

    return docx_path, pdf_path


def generate_cover_letter_from_file(content_path: str | Path, output_folder: str | Path = None) -> tuple[Path, Path | None]:
    content_path = Path(content_path)
    with open(content_path, 'r', encoding='utf-8') as f:
        content = json.load(f)
    if output_folder is None:
        output_folder = content_path.parent
    return generate_cover_letter(content, output_folder)
